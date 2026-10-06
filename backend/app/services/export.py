import io
import os
from pathlib import Path
from xml.sax.saxutils import escape

from ..schemas import Notes
from .notes import notes_markdown
from .storage import timestamp


def sections(meeting):
    notes = Notes.model_validate(meeting['notes'] or {})
    return [
        ('Summary', [notes.summary or 'Notes have not been generated.']),
        ('Key Discussion Points', notes.key_points),
        ('Decisions', notes.decisions),
        ('Action Items', [f'{"[Done] " if a.completed else ""}{a.task}\nOwner: {a.owner}\nDeadline: {a.deadline}'
                          for a in notes.action_items]),
        ('Open Questions', notes.open_questions),
        ('Important Dates', notes.important_dates),
        ('Transcript', meeting['transcript_text'].splitlines()),
    ]


def export(meeting, format):
    title = meeting['title']
    metadata = f'Date: {meeting["date"]}\nDuration: {timestamp(meeting["duration"])} (recorded audio)'
    if format == 'md':
        notes = notes_markdown(Notes.model_validate(meeting['notes'] or {}))
        content = f'# {title}\n\n{metadata}\n\n{notes}\n## Transcript\n\n{meeting["transcript_text"]}\n'
        return content.encode('utf-8'), 'text/markdown; charset=utf-8'
    if format == 'txt':
        parts = [title, metadata]
        for heading, entries in sections(meeting):
            parts += [heading, '\n'.join(entries) or 'None recorded.']
        return '\n\n'.join(parts).encode('utf-8'), 'text/plain; charset=utf-8'
    output = io.BytesIO()
    if format == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt
        document = Document()
        document.sections[0].top_margin = Inches(0.8)
        document.styles['Normal'].font.name = 'Calibri'
        document.styles['Normal'].font.size = Pt(11)
        document.add_heading(title, 0)
        document.add_paragraph(metadata)
        for heading, entries in sections(meeting):
            document.add_heading(heading, 1)
            for entry in entries or ['None recorded.']:
                document.add_paragraph(entry)
        document.save(output)
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if format == 'pdf':
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.enums import TA_LEFT
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.units import inch
        styles = getSampleStyleSheet()
        font_paths = [os.getenv('MEETNOTES_PDF_FONT', ''), 'C:/Windows/Fonts/arial.ttf',
                      '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf']
        font_path = next((p for p in font_paths if p and Path(p).is_file()), None)
        if font_path:
            if 'MeetNotes' not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont('MeetNotes', font_path))
            for style in styles.byName.values():
                style.fontName = 'MeetNotes'
        styles['Normal'].leading = 16
        styles['Title'].alignment = TA_LEFT
        styles['Heading2'].textColor = colors.HexColor('#263d38')
        def paragraph(text, style):
            return Paragraph(escape(text).replace('\n', '<br/>'), styles[style])
        story = [paragraph(title, 'Title'), paragraph(metadata, 'Normal'), Spacer(1, 0.2 * inch)]
        for heading, entries in sections(meeting):
            story.append(paragraph(heading, 'Heading2'))
            for entry in entries or ['None recorded.']:
                story += [paragraph(entry, 'Normal'), Spacer(1, 6)]
        def footer(canvas, doc):
            canvas.setFont('Helvetica', 8)
            canvas.drawRightString(7.5 * inch, 0.4 * inch, f'MeetNotes · {doc.page}')
        SimpleDocTemplate(output, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54).build(
            story, onFirstPage=footer, onLaterPages=footer)
        return output.getvalue(), 'application/pdf'
    raise ValueError('Unsupported export format')
